"""Private, staff-facing QR check-in web application."""
import hmac
import os
import secrets
from datetime import datetime, timezone
from functools import wraps
from io import BytesIO

from flask import (
    Flask,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from waitress import serve

import config
import db
from qr_utils import verify_token
from texts import MODES, ROLES


def conference_day():
    return datetime.now(config.TIMEZONE).date().isoformat()


app = Flask(__name__)
app.secret_key = f"{config.QR_SECRET}:check-in-session"
app.config.update(
    MAX_CONTENT_LENGTH=4096,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=config.CHECKIN_COOKIE_SECURE,
)


@app.after_request
def prevent_private_page_caching(response):
    response.headers["Cache-Control"] = "no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    return response


@app.get("/health")
def health():
    return jsonify(status="ok")


def staff_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("staff_authenticated") is not True:
            if request.path.startswith("/api/"):
                return jsonify(error="Staff login required."), 401
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def require_csrf():
    expected = session.get("csrf_token", "")
    supplied = request.headers.get("X-CSRF-Token") or request.form.get("csrf_token", "")
    if not expected or not hmac.compare_digest(expected, supplied):
        abort(400, description="Invalid security token. Reload the page and try again.")


def read_qr_token():
    payload = request.get_json()
    token = payload.get("token") if isinstance(payload, dict) else None
    if not isinstance(token, str) or not token.strip() or len(token) > 256:
        return None

    token = token.strip()
    reg_code = verify_token(token)
    if not reg_code:
        return None

    today = conference_day()
    attendee = db.get_registration_by_code(reg_code, today)
    if (
        attendee is None
        or attendee["status"] != "confirmed"
        or not hmac.compare_digest(attendee["qr_token"] or "", token)
    ):
        return None
    return reg_code, attendee, today


def attendee_json(attendee, attendance_date):
    return {
        "registration_code": attendee["reg_code"],
        "name": attendee["full_name"],
        "profession": attendee["role"],
        "institution": attendee["institution"],
        "attendance_mode": attendee["attendance_mode"],
        "attendance_date": attendance_date,
        "checked_in_at": attendee["checked_in_at"],
    }


def export_label(options, value):
    for key, label, _amharic_label in options:
        if key == value:
            return label.split(maxsplit=1)[-1]
    return str(value).replace("_", " ").title()


@app.route("/login", methods=["GET", "POST"])
def login():
    if not config.CHECKIN_PASSWORD:
        abort(503, description="Set CHECKIN_PASSWORD in the project .env file first.")

    if request.method == "GET":
        session.setdefault("csrf_token", secrets.token_urlsafe(32))
        return render_template("checkin_login.html", csrf_token=session["csrf_token"])

    supplied_csrf = request.form.get("csrf_token", "")
    if not hmac.compare_digest(session.get("csrf_token", ""), supplied_csrf):
        abort(400, description="Invalid security token. Reload the page and try again.")

    password = request.form.get("password", "")
    if not hmac.compare_digest(config.CHECKIN_PASSWORD, password):
        flash("Incorrect staff password. Please try again.", "error")
        return redirect(url_for("login"))

    session.clear()
    session["staff_authenticated"] = True
    session["csrf_token"] = secrets.token_urlsafe(32)
    return redirect(url_for("index"))


@app.post("/logout")
@staff_required
def logout():
    require_csrf()
    session.clear()
    return redirect(url_for("login"))


@app.get("/")
@staff_required
def index():
    return render_template("checkin.html", csrf_token=session["csrf_token"])


@app.get("/export/registrations.xlsx")
@staff_required
def export_registrations():
    headers = [
        "Timestamp",
        "Title",
        "Full name, including your grandfather’s name.",
        "Profession",
        "Organization",
        "Email",
        "Phone number",
        "Which conference day will you attend",
        "How did you hear about this event",
        "Region",
        "Are you a student, intern, resident, or fellow",
        "Registration code",
        "Registration status",
        "Attendance type",
    ]
    registrations = db.get_registrations_for_export()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Registrations"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.sheet_properties.tabColor = "145C4C"
    workbook.properties.title = "ASOC 2026 Registration List"
    workbook.properties.subject = "Conference attendee registrations"

    header_fill = PatternFill("solid", fgColor="145C4C")
    header_font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
    header_border = Border(
        bottom=Side(style="medium", color="D9AD58")
    )
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = header_border
    worksheet.row_dimensions[1].height = 42

    widths = [21, 14, 36, 24, 30, 34, 20, 38, 34, 20, 36, 20, 20, 20]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width

    for row_number, registration in enumerate(registrations, start=2):
        created_at = datetime.fromisoformat(registration["created_at"])
        if created_at.tzinfo is None:
            created_at = created_at.replace(tzinfo=timezone.utc)
        local_timestamp = created_at.astimezone(config.TIMEZONE).replace(tzinfo=None)
        values = [
            local_timestamp,
            "",
            registration["full_name"],
            export_label(ROLES, registration["role"]),
            registration["institution"],
            registration["email"],
            registration["phone"],
            "",
            "",
            "",
            "",
            registration["reg_code"],
            str(registration["status"]).replace("_", " ").title(),
            export_label(MODES, registration["attendance_mode"]),
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column == 1:
                cell.number_format = "yyyy-mm-dd hh:mm"
            else:
                cell.number_format = "@"
                cell.data_type = "s"
        if row_number % 2 == 0:
            for cell in worksheet[row_number]:
                cell.fill = PatternFill("solid", fgColor="F3F7F5")

    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(registrations) + 1}"
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename_date = datetime.now(config.TIMEZONE).strftime("%Y%m%d")
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"ASOC-2026-registrations-{filename_date}.xlsx",
    )


@app.post("/api/lookup")
@staff_required
def lookup():
    require_csrf()
    result = read_qr_token()
    if result is None:
        return jsonify(error="QR code is invalid or not registered."), 404

    _reg_code, attendee, today = result
    return jsonify(attendee_json(attendee, today))


@app.post("/api/checkin")
@staff_required
def checkin():
    require_csrf()
    result = read_qr_token()
    if result is None:
        return jsonify(error="QR code is invalid or not registered."), 404

    reg_code, _attendee, today = result
    attendee, newly_checked_in = db.check_in_registration(reg_code, today)
    if attendee is None:
        return jsonify(error="Registration could not be found."), 404
    if attendee["status"] != "confirmed":
        return jsonify(error="This registration is not confirmed."), 409

    response = attendee_json(attendee, today)
    response.update(
        message="Checked in successfully." if newly_checked_in else "Already checked in.",
        already_checked_in=not newly_checked_in,
    )
    return jsonify(response)


if __name__ == "__main__":
    if not config.CHECKIN_PASSWORD:
        raise SystemExit(
            "Set CHECKIN_PASSWORD in the project .env file before starting check-in."
        )
    if len(config.CHECKIN_PASSWORD) < 16:
        raise SystemExit(
            "CHECKIN_PASSWORD must be at least 16 characters. Generate a random password."
        )
    if not config.QR_SECRET:
        raise SystemExit(
            "Set QR_SECRET in the project .env file before starting check-in."
        )
    db.init_db()
    port = int(os.getenv("PORT", "8080"))
    host = os.getenv("CHECKIN_HOST", "127.0.0.1")
    serve(app, host=host, port=port)
