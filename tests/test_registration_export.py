import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

import checkin_app
import schedule_handlers
from texts import t


class RegistrationExportTests(unittest.TestCase):
    def setUp(self):
        self.client = checkin_app.app.test_client()

    def authenticate(self):
        with self.client.session_transaction() as session:
            session["staff_authenticated"] = True
            session["csrf_token"] = "test-csrf"

    def test_export_requires_staff_login(self):
        response = self.client.get("/export/registrations.xlsx")

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.headers["Location"], "/login")

    def test_export_downloads_formatted_attendee_workbook(self):
        self.authenticate()
        registration = {
            "created_at": "2026-10-14 12:59:13",
            "reg_code": "ASOC-0001",
            "status": "confirmed",
            "full_name": "=2+2",
            "role": "doctor",
            "institution": "Addis Ababa University",
            "email": "attendee@example.org",
            "phone": "+251963333388",
            "attendance_mode": "in_person",
        }
        with patch.object(
            checkin_app.db,
            "get_registrations_for_export",
            return_value=[registration],
        ):
            response = self.client.get("/export/registrations.xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("ASOC-2026-registrations-", response.headers["Content-Disposition"])
        self.assertEqual(response.headers["Cache-Control"], "no-store")

        workbook = load_workbook(BytesIO(response.data), data_only=False)
        worksheet = workbook["Registrations"]
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, "A1:N2")
        self.assertEqual(
            [worksheet.cell(1, col).value for col in range(1, 12)],
            [
                "Timestamp",
                "Title",
                "Full name, including your grandfather’s name.",
                "Profession",
                "Organization",
                "Email",
                "Phone number",
                "Which conference day will you attend",
                "How did you hear about this event",
                "Region",
                "Are you a student, intern, resident, or fellow",
            ],
        )
        self.assertEqual(worksheet["A2"].value.isoformat(), "2026-10-14T15:59:13")
        self.assertEqual(worksheet["C2"].value, "=2+2")
        self.assertEqual(worksheet["C2"].data_type, "s")
        self.assertEqual(worksheet["E2"].value, "Addis Ababa University")
        self.assertEqual(worksheet["G2"].value, "+251963333388")
        self.assertIsNone(worksheet["B2"].value)
        self.assertIsNone(worksheet["H2"].value)
        self.assertEqual(worksheet["L2"].value, "ASOC-0001")
        self.assertEqual(worksheet["D2"].value, "Doctor")
        self.assertEqual(worksheet["N2"].value, "In person")
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertTrue(worksheet["A1"].fill.fgColor.rgb.endswith("145C4C"))

    def test_staff_home_page_offers_excel_export(self):
        self.authenticate()

        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Download Excel report", response.data)
        self.assertIn(b"/export/registrations.xlsx", response.data)

    def test_schedule_keyboards_offer_back_to_main_menu(self):
        day_list_keyboard = schedule_handlers.days_keyboard("en")
        schedule_day_keyboard = schedule_handlers.day_keyboard(
            "en", 0, "all", 0, 1, ["main"]
        )

        self.assertEqual(
            day_list_keyboard.inline_keyboard[-1][0].callback_data, "menu:home"
        )
        self.assertEqual(
            schedule_day_keyboard.inline_keyboard[-1][0].callback_data, "menu:home"
        )

    def test_contact_details_are_available_in_both_languages(self):
        for language in ("en", "am"):
            contact = t(language, "contact_info")
            self.assertIn("ethiopiastrokefoundation@gmail.com", contact)
            self.assertIn("+251963333388", contact)
            self.assertIn("+251955235523", contact)

    def test_database_export_query_returns_registration_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "test.db"
            with patch.object(checkin_app.db.config, "DB_PATH", str(database_path)):
                checkin_app.db.init_db()
                checkin_app.db.create_registration(
                    123,
                    {
                        "full_name": "Test Attendee",
                        "phone": "+251963333388",
                        "email": "test@example.org",
                        "role": "doctor",
                        "institution": "Test Hospital",
                        "language": "en",
                        "attendance_mode": "in_person",
                    },
                )

                registrations = checkin_app.db.get_registrations_for_export()

        self.assertEqual(len(registrations), 1)
        self.assertEqual(registrations[0]["full_name"], "Test Attendee")
        self.assertEqual(registrations[0]["institution"], "Test Hospital")
        self.assertEqual(registrations[0]["attendance_mode"], "in_person")


if __name__ == "__main__":
    unittest.main()
